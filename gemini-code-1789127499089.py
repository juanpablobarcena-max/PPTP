import streamlit as st
from docx import Document
import io
import re

# Configuración de la página web
st.set_page_config(page_title="Generador de Pliegos", page_icon="📄", layout="centered")

def obtener_numeracion(texto):
    """Busca si el párrafo empieza por un Capítulo o un número tipo 1.4, 1.15.2, etc."""
    # Busca patrones como "CAPITULO II", "1", "1.4", "1.15.2"
    match = re.match(r'^(CAP[IÍ]TULO \w+|1(?:\.\d+)*)', texto.upper())
    if match:
        return match.group(1)
    return None

def eliminar_parrafo(parrafo):
    """Elimina el párrafo del XML sin romper el documento"""
    p = parrafo._element
    p.getparent().remove(p)
    parrafo._p = parrafo._element = None

def main():
    st.title("Generador de Pliegos a Medida 📄")
    st.markdown("Sube tu documento matriz (`.docx`). El sistema detectará los apartados y podrás generar un pliego limpio con el formato original intacto.")
    
    # 1. Subida del archivo
    archivo_subido = st.file_uploader("Sube el documento madre (.docx)", type=["docx"])
    
    if archivo_subido is not None:
        try:
            doc = Document(archivo_subido)
            
            # 2. Escaneo del documento para extraer los títulos
            secciones_variables = {}
            for p in doc.paragraphs:
                texto = p.text.strip()
                if texto:
                    sec = obtener_numeracion(texto)
                    if sec and sec.startswith("1."):
                        partes = sec.split(".")
                        if len(partes) >= 2 and partes[1].isdigit():
                            num = int(partes[1])
                            if 4 <= num <= 52: # Filtramos del 1.4 al 1.52
                                base_sec = f"1.{num}"
                                if base_sec not in secciones_variables:
                                    # Guardamos el título para mostrarlo en pantalla
                                    secciones_variables[base_sec] = texto[:85] + ("..." if len(texto) > 85 else "")
            
            if not secciones_variables:
                st.warning("⚠️ No se han detectado los apartados. Asegúrate de que los números (1.4, 1.5...) están escritos en el texto y no como una lista automática de Word.")
                return

            st.success("✅ Documento analizado correctamente.")
            st.write("---")
            st.subheader("Selecciona los apartados a MANTENER:")
            st.info("💡 Recuerda: Los puntos del 1.1 al 1.3.2 se incluirán vacíos de forma automática.")
            
            # 3. Crear las casillas de verificación
            selecciones = []
            for sec, titulo in secciones_variables.items():
                if st.checkbox(titulo, value=True): # Vienen marcados por defecto
                    selecciones.append(sec)
                    
            st.write("---")
            
            # 4. Botón de generar
            if st.button("🚀 Generar Nuevo Documento", type="primary"):
                with st.spinner("Procesando el documento, por favor espera..."):
                    nuevo_doc = procesar_documento(Document(archivo_subido), selecciones)
                    
                    # Guardar el documento en la memoria temporal del servidor
                    buffer = io.BytesIO()
                    nuevo_doc.save(buffer)
                    buffer.seek(0)
                    
                st.balloons()
                st.success("¡Tu nuevo pliego a medida está listo!")
                
                # 5. Botón de descarga
                st.download_button(
                    label="📥 Descargar Pliego Generado",
                    data=buffer,
                    file_name="Nuevo_Pliego_A_Medida.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
                
        except Exception as e:
            st.error(f"Se ha producido un error al leer el archivo: {e}")

def procesar_documento(doc, selecciones):
    fijos = ["1", "1.1", "1.1.1", "1.1.2", "1.1.3", "1.2", "1.3", "1.3.1", "1.3.2"]
    estado = "MANTENER"
    
    for p in doc.paragraphs:
        texto = p.text.strip()
        es_cabecera = False
        
        if texto:
            sec = obtener_numeracion(texto)
            if sec:
                if sec.startswith("CAP"):
                    estado = "MANTENER"
                    es_cabecera = True
                elif sec in fijos:
                    estado = "VACIAR"
                    es_cabecera = True
                elif sec.startswith("1."):
                    partes = sec.split(".")
                    if len(partes) >= 2 and partes[1].isdigit():
                        base_sec = f"1.{partes[1]}"
                        if base_sec in selecciones:
                            estado = "MANTENER"
                        else:
                            estado = "BORRAR"
                        es_cabecera = True
        
        # Aplicamos la acción correspondiente (VACIAR o BORRAR)
        if estado == "VACIAR":
            # Si es vaciar, borramos el contenido pero respetamos el título de cabecera
            if not es_cabecera and texto != "":
                eliminar_parrafo(p)
        elif estado == "BORRAR":
            # Si es borrar, nos cargamos el párrafo entero (título incluido)
            eliminar_parrafo(p)
            
    return doc

if __name__ == "__main__":
    main()